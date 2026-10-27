import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_components_excel_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Components Template"

    headers = [
        "Name", "Category", "Quantity", "Unit", "Min Alert Qty", 
        "Rack Location", "Specs", "Invoice / Bill No", "Purchase Date", "Vendor Name", "Unit Price"
    ]
    ws.append(headers)

    sample_rows = [
        ["RJ45 Connectors", "Consumable", 100, "Nos", 20, "Rack A-1", "Cat6 Crimping Connectors", "INV-2026-901", "2026-08-15", "D-Link Distributor", 8.50],
        ["Cat6 LAN Cable Roll", "Consumable", 50, "Meters", 15, "Rack A-2", "D-Link Ethernet Cable", "INV-2026-902", "2026-08-15", "D-Link Distributor", 25.00],
        ["Dell USB Mouse", "Peripherals/Hardware", 20, "Nos", 5, "Cupboard 1", "Wired Optical Mouse", "Opening Stock", "2025-01-10", "Existing Lab Stock", 250.00],
        ["HP Standard Keyboard", "Peripherals/Hardware", 15, "Nos", 4, "Cupboard 1", "USB Multimedia Keyboard", "DSR Page 12", "2024-11-05", "College Central Store", 450.00],
        ["CMOS Battery (CR2032)", "Consumable", 10, "Nos", 5, "Drawer B-1", "3V Lithium Cell", "CASH-BILL-88", "2026-09-01", "Shinde Electronics", 30.00],
        ["Crimping Tool Kit", "Tools", 5, "Nos", 2, "Toolbox 1", "RJ45/RJ11 Crimping & Punch Tool", "Opening Stock", "2024-06-20", "Existing Lab Stock", 600.00]
    ]
    for row in sample_rows:
        ws.append(row)

    # Style Header
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_students_excel_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students Template"

    headers = ["UID No", "Full Name", "Department", "Year", "Division", "Phone"]
    ws.append(headers)

    sample_rows = [
        ["UID-2024-001", "Rahul Sharma", "Computer Engineering", "TE", "A", "9876543210"],
        ["UID-2024-002", "Sneha Patil", "Computer Engineering", "TE", "B", "9876543211"],
        ["UID-2024-011", "Amit Deshmukh", "Information Technology", "BE", "A", "9876543212"],
        ["UID-2024-018", "Pooja Jadhav", "AI & Data Science", "SE", "A", "9876543213"]
    ]
    for row in sample_rows:
        ws.append(row)

    # Style Header
    header_fill = PatternFill(start_color="1E7145", end_color="1E7145", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def parse_components_file(file_storage):
    filename = file_storage.filename.lower()
    records = []

    if filename.endswith('.xlsx'):
        wb = openpyxl.load_workbook(file_storage, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return records
        
        for row in rows[1:]:
            if not row or not any(row):
                continue
            name = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
            if not name:
                continue
            category = str(row[1]).strip() if len(row) > 1 and row[1] is not None else "Consumable"
            try:
                quantity = int(row[2]) if len(row) > 2 and row[2] is not None else 0
            except:
                quantity = 0
            unit = str(row[3]).strip() if len(row) > 3 and row[3] is not None else "Nos"
            try:
                min_alert = int(row[4]) if len(row) > 4 and row[4] is not None else 5
            except:
                min_alert = 5
            rack = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
            specs = str(row[6]).strip() if len(row) > 6 and row[6] is not None else ""
            
            invoice_no = str(row[7]).strip() if len(row) > 7 and row[7] is not None else "Opening Stock"
            purchase_date = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""
            vendor_name = str(row[9]).strip() if len(row) > 9 and row[9] is not None else ""
            try:
                unit_price = float(row[10]) if len(row) > 10 and row[10] is not None else 0.0
            except:
                unit_price = 0.0

            records.append({
                'name': name,
                'category': category,
                'quantity': quantity,
                'unit': unit,
                'min_alert_qty': min_alert,
                'rack_location': rack,
                'specs': specs,
                'invoice_no': invoice_no,
                'purchase_date': purchase_date,
                'vendor_name': vendor_name,
                'unit_price': unit_price
            })

    elif filename.endswith('.csv'):
        content = file_storage.stream.read().decode('utf-8-sig', errors='ignore')
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        if not rows:
            return records
        
        for row in rows[1:]:
            if not row or not any(row):
                continue
            name = row[0].strip() if len(row) > 0 else ""
            if not name:
                continue
            category = row[1].strip() if len(row) > 1 and row[1].strip() else "Consumable"
            try:
                quantity = int(row[2].strip()) if len(row) > 2 else 0
            except:
                quantity = 0
            unit = row[3].strip() if len(row) > 3 and row[3].strip() else "Nos"
            try:
                min_alert = int(row[4].strip()) if len(row) > 4 else 5
            except:
                min_alert = 5
            rack = row[5].strip() if len(row) > 5 else ""
            specs = row[6].strip() if len(row) > 6 else ""
            invoice_no = row[7].strip() if len(row) > 7 and row[7].strip() else "Opening Stock"
            purchase_date = row[8].strip() if len(row) > 8 else ""
            vendor_name = row[9].strip() if len(row) > 9 else ""
            try:
                unit_price = float(row[10].strip()) if len(row) > 10 else 0.0
            except:
                unit_price = 0.0

            records.append({
                'name': name,
                'category': category,
                'quantity': quantity,
                'unit': unit,
                'min_alert_qty': min_alert,
                'rack_location': rack,
                'specs': specs,
                'invoice_no': invoice_no,
                'purchase_date': purchase_date,
                'vendor_name': vendor_name,
                'unit_price': unit_price
            })

    return records


def parse_students_file(file_storage):
    filename = file_storage.filename.lower()
    records = []

    if filename.endswith('.xlsx'):
        wb = openpyxl.load_workbook(file_storage, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return records
        
        for row in rows[1:]:
            if not row or not any(row):
                continue
            uid_no = str(row[0]).strip().upper() if len(row) > 0 and row[0] is not None else ""
            name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            if not uid_no or not name:
                continue
            dept = str(row[2]).strip() if len(row) > 2 and row[2] is not None else "Computer Engineering"
            year = str(row[3]).strip().upper() if len(row) > 3 and row[3] is not None else "TE"
            division = str(row[4]).strip().upper() if len(row) > 4 and row[4] is not None else "A"
            phone = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
            academic_year = str(row[6]).strip() if len(row) > 6 and row[6] is not None else "2026-2027"

            records.append({
                'roll_no': uid_no, # mapped to student identifier
                'name': name,
                'department': dept,
                'year': year,
                'division': division,
                'phone': phone,
                'academic_year': academic_year or "2026-2027"
            })

    elif filename.endswith('.csv'):
        content = file_storage.stream.read().decode('utf-8-sig', errors='ignore')
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        if not rows:
            return records
        
        for row in rows[1:]:
            if not row or not any(row):
                continue
            uid_no = row[0].strip().upper() if len(row) > 0 else ""
            name = row[1].strip() if len(row) > 1 else ""
            if not uid_no or not name:
                continue
            dept = row[2].strip() if len(row) > 2 and row[2].strip() else "Computer Engineering"
            year = row[3].strip().upper() if len(row) > 3 and row[3].strip() else "TE"
            division = row[4].strip().upper() if len(row) > 4 else "A"
            phone = row[5].strip() if len(row) > 5 else ""
            academic_year = row[6].strip() if len(row) > 6 and row[6].strip() else "2026-2027"

            records.append({
                'roll_no': uid_no,
                'name': name,
                'department': dept,
                'year': year,
                'division': division,
                'phone': phone,
                'academic_year': academic_year or "2026-2027"
            })

    return records


def export_transactions_to_excel(transactions_list, lab_name="All Laboratories"):
    """
    Generate a beautifully styled Excel workbook containing complete transaction records
    suitable for official college lab audits, NAAC/NBA verification, and inventory registers.
    """
    from datetime import datetime

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions Register"

    # Title Banner Block
    ws.merge_cells('A1:P1')
    ws['A1'] = "ST. VINCENT PALLOTTI COLLEGE OF ENGINEERING & TECHNOLOGY, NAGPUR"
    ws['A1'].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A2:P2')
    ws['A2'] = f"Department of Industrial IoT — {lab_name} | Equipment Issue & Return Register"
    ws['A2'].font = Font(name="Calibri", size=11, bold=True, color="404040")
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A3:P3')
    ws['A3'] = f"Generated On: {datetime.now().strftime('%d-%m-%Y %I:%M %p')} | Total Records: {len(transactions_list)}"
    ws['A3'].font = Font(name="Calibri", size=10, italic=True, color="7F7F7F")
    ws['A3'].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        "Trans ID", "Date & Time", "Laboratory", "Recipient Type",
        "Recipient Name", "UID / Roll / ID", "Department", "Contact Phone",
        "Component Name", "Qty", "Unit", "Purpose",
        "Status", "Expected Return", "Actual Return", "Return Remarks / Issues"
    ]
    ws.append([]) # Empty row 4
    ws.append(headers) # Row 5

    header_row = 5
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    alt_fill = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")

    for idx, t in enumerate(transactions_list, start=1):
        def _get(key, fallback=""):
            try:
                v = t[key]
                return v if v is not None else fallback
            except Exception:
                return fallback

        r_type = _get('recipient_type') or _get('display_type') or 'STUDENT'
        r_name = _get('recipient_name') or _get('display_name') or 'N/A'
        r_code = _get('recipient_id_code') or _get('display_id_code') or ''
        r_dept = _get('recipient_dept') or _get('display_dept') or ''
        r_phone = _get('recipient_phone') or _get('display_phone') or ''
        c_name = _get('component_name') or ''
        qty = _get('quantity', 0)
        unit = _get('unit', 'Nos')
        purpose = _get('purpose', '')
        status = _get('status', 'ISSUED')
        exp_date = _get('expected_return_date', '')
        ret_date = _get('return_date', '')
        remarks = _get('damage_remarks') or _get('remarks') or ''
        
        dmg = _get('returned_damaged_qty', 0)
        mis = _get('returned_missing_qty', 0)
        if dmg or mis:
            issues = []
            if dmg: issues.append(f"Damaged: {dmg}")
            if mis: issues.append(f"Missing: {mis}")
            remarks = f"[{', '.join(issues)}] {remarks}".strip()

        lab = _get('lab_name', lab_name)
        dt = _get('transaction_date', '')

        row_data = [
            _get('id'), dt, lab, r_type,
            r_name, r_code, r_dept, r_phone,
            c_name, qty, unit, purpose,
            status, exp_date, ret_date, remarks
        ]
        ws.append(row_data)
        current_row = ws.max_row
        for col_num in range(1, len(row_data) + 1):
            c = ws.cell(row=current_row, column=col_num)
            c.border = thin_border
            if idx % 2 == 0:
                c.fill = alt_fill
            c.alignment = Alignment(vertical="center")
            if col_num in [1, 10, 14, 15]:
                c.alignment = Alignment(horizontal="center", vertical="center")

    # Column Auto-Width
    for col in ws.columns:
        if col[0].row > 3:
            max_len = max(len(str(cell.value or '')) for cell in col if cell.row >= header_row)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, 38), 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
