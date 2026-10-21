import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import shutil
import os

# -------------------------------------------------------------
# 1. STUDENTS DATA (All under Academic Year 2026-2027)
# Department of Industrial IoT - St. Vincent Pallotti College of Engg, Nagpur
# -------------------------------------------------------------

# Sem 5 (Academic Year 2026-2027) - 71 Students
sem5_raw = [
    ("24009001", "TILAK PURUSHOTTAM DHAPODKAR", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009002", "MANTHAN SUNIL SHENDE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009003", "OJAL GAJANAN SHAHARE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009004", "ROHAN RAVINDRA SHEMBEKAR", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009005", "ADITYA PRASHANT FALKE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009006", "VEDANT PADMAKAR MANKAR", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009008", "SANYUKTA SUNIL GADE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009009", "PULKIT MAHESH WATH", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009010", "SHANTANU BAPURAO BHURE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009011", "LALIT ASHOK RATHOD", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009012", "LAVANYA SUNIL CHOUDHARY", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009013", "GAURAV RAJENDRA SOMKUWAR", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009014", "YOGESHWAR ANANTRAO BAHADURE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009015", "LEENA VIJAY RAHANGDALE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009016", "KHUSHBU PRAMOD NAVGHARE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009017", "TEJAS ASHOK MULE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009018", "ROMIT RAJENDRA TABHANE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009019", "SUHANI RAMBHAU MAHAJAN", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009020", "HARSHAL RAJESH NANDANWAR", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009021", "SALONI DNYANESHWAR CHIPDE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009022", "BHUVNESHWARI DIGAMBER CHARDE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009023", "ANJALI ANIL HARODE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009024", "CHAITRALI PRAVIN BELE", "TE (Sem V)", "Batch-1", "2026-2027"),
    ("24009025", "KSHITIJ RAJU ADAKANE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009026", "CHYARUL DASHARATH SAWAJKAR", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009027", "ADITYA LAXMAN DHANVIJ", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009028", "KHUSHI RAMESH NEWARE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009029", "ATHARVA JITENDRA KALODE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009030", "KHUSHI HARSHUNANDAN BARDE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009031", "VIDHI VILAS UKEY", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009032", "SANJHI ROSHAN SADRANI", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009033", "VEDANT PRAVIN CHAMAT", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009035", "KASHISH KAILASH MUNESHWAR", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009036", "SUJAL RAM NAGOSE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009037", "SONIYA SHANKAR HARODE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009038", "APURVA UPENDRA MOHITE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009040", "AACHAL DIWAKAR BORATE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009041", "ROOPAM YASHWANT ZADE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009042", "AYUSH AMBADAS NARNAWARE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009043", "PRADNYA CHANDRASHEKHAR KHONDE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009044", "SUMEDH MILIND NIRWAN", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009045", "DIVYA DHARMENDRA SHARNAGAT", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009046", "ARPIT VIVEK KHARWADE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009047", "ARYA KHILESH CHOUDHARY", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009048", "VIVEK SHEKHAR DAKHANE", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009049", "JOSBIN T BIJU", "TE (Sem V)", "Batch-2", "2026-2027"),
    ("24009050", "LAVANSH ANIL NAIR", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009051", "R BHAVYA R MANISH NAIDU", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009052", "PUNIT TUKARAM NIMBALKAR", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009053", "ANSHUMAN VILAS PATIL", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009054", "SHRAVAN RAJESH SAKHARKAR", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009055", "DHANASHREE VIVEK SOITKAR", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009056", "TANISHKA MILIND TAPAS", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009057", "VEDANT VIJAY AMBULKAR", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009058", "ABHAYA JAYANT DEOGADE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009059", "DEVESH SHYAMRAO SAMBARE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009060", "RISHANK ANIL DHOBE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009062", "LALCHAND MANOHAR KUHIKAR", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009063", "HEMANT MITARAM PARDHI", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009064", "RAMANPREET BALDEVSINGH KAUR", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009065", "PRIYANSHU RANJEET SINGH", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("24009066", "YASH PRADEEP WANDHARE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("23017061", "AYUSHI ASHISH NAGRARE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("22017028", "MRUNAL SAWARKAR", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("25109001", "KHUSHI TULSIRAM THAKRE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("25109002", "YUG VIKRANT RAMTEKE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("25109003", "PRACHI DIWAKAR GOHATE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("25109004", "VAIBHAV JIVAN GHODE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("25109005", "SANIKA JAGDISH HARDE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("25109006", "TUSHAR PRAMOD AGASHE", "TE (Sem V)", "Batch-3", "2026-2027"),
    ("25109007", "PARTH MANOJ JANAI", "TE (Sem V)", "Batch-3", "2026-2027"),
]

# Sem 6 (Academic Year 2026-2027) - 72 Students
sem6_raw = [
    ("22017021", "ARJUN DABHOLKAR", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("22017033", "YASH CHIMANTRAWAR", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("22017041", "SRUSHTI BHAGWAT", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("22017061", "SUBHOJEET GUPTA", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017001", "SALONI UMESH GAWATE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017002", "BHAGYASHREE HARICHANDRA DHAKATE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017003", "SHRUTI BHASKAR WAHANE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017004", "ANUSHKA DNYANESHWAR BHOPE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017005", "KUNAL PRAVIN SHELKE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017006", "MOHIT PRAMOD UKUDDE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017007", "BHUMIKA BHANUPRATAP DEOTARE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017009", "VANASHREE NILESH BAREKAR", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017010", "ISHA SHIVKUMAR BAGHELE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017012", "RUGVED NARENDRA PARDHY", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017013", "TEJAS SANJAY DUNEDAR", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017014", "SWATI TILAKCHAND CHHANIKAR", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017015", "PRATHAM VINOD CHAKOLE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017016", "ARTHIK NASHIK CHAWARE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017017", "BHUMIKA VIKAS SONTAKKE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017018", "PARAG MANOJ BHAGAT", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017019", "TRUPTI PRAMOD THAWKAR", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017020", "SAANCHI BHARAT KAMBLE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017021", "TANUSHREE SANJAY NAGESHWAR", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017023", "ACHAL AMOL DIWASE", "TE (Sem VI)", "Batch-1", "2026-2027"),
    ("23017024", "SHRAVANI SUBHASH THAKRE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017026", "DHANASHRI RAVINDRA MATE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017027", "SUVIDHA VIJAY KANTODE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017028", "SANSKRUTI MANOJ BINEKAR", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017029", "SHARVARI GYANDEO PACHARE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017030", "JANVI VILAS BHOYAR", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017031", "CHINMAI RAMESH MALAPOLU", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017032", "ANSHUL PRAVIN KAKDE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017033", "SARANSHI SUNIL NANDANWAR", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017034", "VED AKASH AMTE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017035", "SAMRUDDHI SARANG NERKAR", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017036", "VEDANTI RAVINDRA BHAYDE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017037", "CHINMAY PRAFUL WADETTIWAR", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017038", "ALPA ANIRUDDHA PANDE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017040", "ARYA RAJENDRA DHOTE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017041", "RIYA RAVINDRA BELEKAR", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017042", "OM PRAVIN MISAL", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017043", "LAXMIKANT HEMRAJ BHOSKAR", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017045", "GAURAV PRAKASH KOSARE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017046", "NIKUNJ SANJAY PATIL", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017047", "TANAY DEVENDRA RAUT", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017048", "SWAYAM PRAKASH KUKDE", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017049", "SHRAVANI DEVIDAS JOSHI", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017050", "SEJAL ASHOK SAHARKAR", "TE (Sem VI)", "Batch-2", "2026-2027"),
    ("23017051", "AVINASH SHRIKRUSHNA WANKHADE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017052", "MAYANK MANISH KENE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017053", "LAUKIK PRAKASH CHIVHANE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017054", "OMKAR RAVINDRA KULKARNI", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017055", "VIBHANSHU UMESH MESHRAM", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017056", "ADITYA SACHIN LIKHITKAR", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017057", "SHUBHAM NITIN KULKARNI", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017058", "AARYAN NILESH VITALKAR", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017059", "DEEPANSHU SUNIL DHAKITE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017060", "RUDRA ATUL DAHIKAR", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017062", "PRANVATI VINAY LANJWAR", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017063", "SHREYA MORESHWAR SATHAWANE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017064", "ABOLI AJAYRAO WATKAR", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017065", "AYUSH VIJAY NIMJE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23017066", "OMSAIL VILAS RAJURKAR", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("23117001", "RUSHIKESH DHOTE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("24109001", "AAMOD KALEY", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("24109002", "ANSH UKEY", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("24109003", "DIVYA BURRIWAR", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("24109004", "ABHISHEK RAUT", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("24109005", "MRUNALI GULHANE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("24109006", "DISHANT MUNGLE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("24109007", "SARVESH GOTMARE", "TE (Sem VI)", "Batch-3", "2026-2027"),
    ("24109008", "ABHIJEET SHINDE", "TE (Sem VI)", "Batch-3", "2026-2027"),
]

# Sem 3 (Academic Year 2026-2027) - 69 Students
sem3_raw = [
    ("25009001", "PRACHI PATIKUMAR SHAHARE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009003", "SWAGAT MANGESH BHURE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009005", "KASHIB KHURSHID ANSARI", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009006", "BHAVIKA PRADIP SATPUTE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009007", "CHIRAYU MAHENDRA MANKAR", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009008", "SOHAM KISHOR JANGADE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009009", "PARI SATISH CHILLURE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009010", "RACHI MAHESH YENURKAR", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009012", "ARPIT VIRENDRA TEMBHEKAR", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009013", "OM SHAILENDRA CHARPE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009014", "OM RAKESH PATIL", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009015", "TANMAY VILAS RAUT", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009016", "BHAVESH BHASKAR ATKARI", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009017", "RAHI EKNATH GHODE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009018", "TEJAS ROSHAN PIMPALGHARE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009019", "SUKHADA MANISH SHINGANE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009020", "LAVINYA CHANDRASHEKHAR KAPGATE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009021", "VIDHANT YOGIRAJ LOKHANDE", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009023", "ADITI SHITAL MOHOD", "SE (Sem III)", "Batch-1", "2026-2027"),
    ("25009024", "AYAN KASHYAP KOSE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009025", "SAMIKSHA UMESHWAR MANMODE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009026", "SANKET SASHIM SAO", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009027", "JANHAVI YOGANAND SONTAKKE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009028", "PRASAD RAJENDRA GAJULWAR", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009029", "ANUSHKA SHARAD AJBALE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009030", "AAYUSH SUNIL WAGHMARE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009031", "NIRMITI MANOJ GADHAVE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009032", "JAYESH VINOD NIMBULKAR", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009033", "ARNAV BHALCHANDRA YELMULE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009034", "PUNIT ASHOK DODKE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009035", "NISHAD RAJESH THAWKAR", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009036", "CHITRANSH SATISH HEDAU", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009037", "SHREYASH PANKAJ PAUNIKAR", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009038", "SRUSHTI SANJAY GHODE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009039", "SUHANI UJAVENDRA BANSOD", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009040", "JANHAVI NITIN NINAVE", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009041", "ANUSHKA VINOD DORLIKAR", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009042", "ARYAN SANGDIL HIREKHAN", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009043", "VEDANT ANILKUMAR BAVASKAR", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009044", "ARYAN KISHOR GHODESWAR", "SE (Sem III)", "Batch-2", "2026-2027"),
    ("25009045", "SHASHWAT ASHWIN JAGDISH", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009047", "TANUSH SUHAS PARATE", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009049", "SARTHAK RAJENDRA BAND", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009050", "ARNAV RAMANPRAKASH NAKTODE", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009051", "DHANSHRI SANJAY RAUT", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009053", "KSHITIJA SHARAD KHAPEKAR", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009054", "PURVA JAGDISH KATWE", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009055", "DIVYA MAHESH KAPSE", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009056", "VRASHI ANIL AKHARE", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009057", "GARGI YOGESH GIRADKAR", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009058", "SUDHANSHU MANGESH YEMBADWAR", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009059", "ADITI HARSHAL KHARKAR", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009060", "ARYAN SURYAPRATAP GUPTA", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009061", "BHAVYA PRASHANT CHIMANKAR", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009062", "KINJAL PRAMOD SARODE", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009063", "SARTHAK SAMIR ANADEO", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009064", "SOHAM SHYAM MAHANKAL", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009065", "DARSHITA VINOD SONKUSARE", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("25009067", "VAIDEHEE SANJAY DESHMUKH", "SE (Sem III)", "Batch-3", "2026-2027"),
    ("26109001", "ADITEE ZOD", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109002", "BHRAMAN PRAMOD BHAGAT", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109003", "CHAITANYA NAGAPURE", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109004", "HIMANSHI SHALIKRAM MESHRAM", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109005", "JUHI SHENDRE", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109006", "KARTIK S KAUNDER", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109007", "OM KAMDE", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109008", "PIYUSH MANOHAR PATLE", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109009", "RUTUJA VILAS CHOBITKAR", "SE (Sem III)", "DSY", "2026-2027"),
    ("26109010", "SAKSHI DHOK", "SE (Sem III)", "DSY", "2026-2027"),
]

# Combined list of all 212 students with Academic Year 2026-2027
all_students = []
for s in sem5_raw:
    all_students.append((s[0], s[1], "Industrial IoT", s[2], s[3], "", s[4]))
for s in sem6_raw:
    all_students.append((s[0], s[1], "Industrial IoT", s[2], s[3], "", s[4]))
for s in sem3_raw:
    all_students.append((s[0], s[1], "Industrial IoT", s[2], s[3], "", s[4]))

all_students.sort(key=lambda x: x[0])

# Styling definition
header_fill_std = PatternFill(start_color="1E7145", end_color="1E7145", fill_type="solid")
header_fill_sub = PatternFill(start_color="2E8B57", end_color="2E8B57", fill_type="solid")
header_font_std = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD')
)

student_headers = ["UID No", "Full Name", "Department", "Year / Semester", "Division / Batch", "Phone", "Academic Year"]

wb_students = openpyxl.Workbook()

# Sheet 1: All Students (AY 2026-2027) - Active Sheet
ws_all = wb_students.active
ws_all.title = "All Students (AY 2026-27)"
ws_all.append(student_headers)

for s in all_students:
    ws_all.append(list(s))

def style_student_sheet(ws, fill_color):
    for cell in ws[1]:
        cell.fill = fill_color
        cell.font = header_font_std
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = border_thin
            if cell.column in [1, 4, 5, 6, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

style_student_sheet(ws_all, header_fill_std)

# Sheet 2: Sem V (AY 2026-2027)
ws_sem5 = wb_students.create_sheet(title="Sem V (AY 2026-27)")
ws_sem5.append(student_headers)
for s in sem5_raw:
    ws_sem5.append([s[0], s[1], "Industrial IoT", s[2], s[3], "", s[4]])
style_student_sheet(ws_sem5, header_fill_sub)

# Sheet 3: Sem VI (AY 2026-2027)
ws_sem6 = wb_students.create_sheet(title="Sem VI (AY 2026-27)")
ws_sem6.append(student_headers)
for s in sem6_raw:
    ws_sem6.append([s[0], s[1], "Industrial IoT", s[2], s[3], "", s[4]])
style_student_sheet(ws_sem6, header_fill_sub)

# Sheet 4: Sem III (AY 2026-2027)
ws_sem3 = wb_students.create_sheet(title="Sem III (AY 2026-27)")
ws_sem3.append(student_headers)
for s in sem3_raw:
    ws_sem3.append([s[0], s[1], "Industrial IoT", s[2], s[3], "", s[4]])
style_student_sheet(ws_sem3, header_fill_sub)

wb_students.save("students_upload.xlsx")
print(f"Saved students_upload.xlsx with 212 students across 4 sheets, all under Academic Year 2026-2027!")

# Copy to artifacts directory
dest_dir = r"C:\Users\student\.gemini\antigravity\brain\93d89031-cddd-4011-9c3f-c27f00e1e50f"
shutil.copy("students_upload.xlsx", os.path.join(dest_dir, "students_upload.xlsx"))
print("Copied updated students_upload.xlsx to artifacts directory!")
